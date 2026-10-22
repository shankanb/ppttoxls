import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def create_sample_deck_standard(output_path: str):
    """Creates a standard corporate-styled budget deck."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Corporate Financial Plan FY2025"
    p1.font.name = "Arial"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    p2 = tf.add_paragraph()
    p2.text = "Executive Budget Review & Departmental Allocations"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # Slide 2: Executive Summary
    s2 = prs.slides.add_slide(blank_layout)
    # Title
    t_box = s2.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11), Inches(1))
    t_frame = t_box.text_frame
    p = t_frame.paragraphs[0]
    p.text = "Executive Summary"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    # Bullets
    b_box = s2.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11), Inches(4))
    b_frame = b_box.text_frame
    bullets = [
        "Fiscal Year: FY2025",
        "Total Approved Budget: $4,500,000 across all core operational divisions",
        "Approved Headcount Target: 85 FTEs by end of Q4",
        "Priority focus on scalable engineering infrastructure and strategic sales expansion",
    ]
    for i, txt in enumerate(bullets):
        bp = b_frame.paragraphs[0] if i == 0 else b_frame.add_paragraph()
        bp.text = f"• {txt}"
        bp.font.size = Pt(18)
        bp.space_after = Pt(14)

    # Slide 3: Department Breakdown Table
    s3 = prs.slides.add_slide(blank_layout)
    t_box3 = s3.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11), Inches(1))
    p3 = t_box3.text_frame.paragraphs[0]
    p3.text = "Departmental Budget Allocation"
    p3.font.size = Pt(28)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    # Table: 5 rows (1 header + 4 depts), 4 cols
    table_shape = s3.shapes.add_table(5, 4, Inches(1), Inches(2.0), Inches(11), Inches(3.5))
    tbl = table_shape.table

    table_data = [
        ["Department", "Planned Spend", "Q1 Actual", "Priority"],
        ["Engineering", "$2,100,000", "$510,000", "Critical"],
        ["Marketing & Sales", "$1,200,000", "$290,000", "High"],
        ["Operations & Support", "$800,000", "$180,000", "Medium"],
        ["Human Resources", "$400,000", "$95,000", "Standard"],
    ]
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            p_cell = cell.text_frame.paragraphs[0]
            p_cell.font.size = Pt(14)
            if r_idx == 0:
                p_cell.font.bold = True
                p_cell.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x2F, 0x55, 0x97)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Saved standard deck: {output_path}")


def create_sample_deck_variant(output_path: str):
    """
    Creates an alternatively styled deck with different fonts, shapes instead of bullets,
    and rearranged table columns, containing identical underlying financial figures.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Variant Title Slide (Darker theme tone)
    s1 = prs.slides.add_slide(blank_layout)
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11), Inches(2.5))
    tf = tb.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "FY2025 Strategic Budget & Allocation"
    p1.font.name = "Georgia"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(0x0F, 0x20, 0x27)

    p2 = tf.add_paragraph()
    p2.text = "Leadership Overview - Financial Targets"
    p2.font.name = "Georgia"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    # Slide 2: Metrics in separate Callout Shapes (NOT bullet list)
    s2 = prs.slides.add_slide(blank_layout)
    t_box = s2.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11), Inches(0.8))
    p = t_box.text_frame.paragraphs[0]
    p.text = "Key Financial Indicators"
    p.font.name = "Georgia"
    p.font.size = Pt(26)
    p.font.bold = True

    # Card 1: Fiscal Year
    card1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.8), Inches(3.4), Inches(2.2))
    card1.fill.solid()
    card1.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF7)
    c1_tf = card1.text_frame
    c1_p1 = c1_tf.paragraphs[0]
    c1_p1.text = "Fiscal Year"
    c1_p1.font.size = Pt(14)
    c1_p1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    c1_p2 = c1_tf.add_paragraph()
    c1_p2.text = "FY2025"
    c1_p2.font.size = Pt(26)
    c1_p2.font.bold = True
    c1_p2.font.color.rgb = RGBColor(0x0F, 0x20, 0x27)

    # Card 2: Total Budget
    card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.8), Inches(3.6), Inches(2.2))
    card2.fill.solid()
    card2.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF7)
    c2_tf = card2.text_frame
    c2_p1 = c2_tf.paragraphs[0]
    c2_p1.text = "Total Authorized Expenditure"
    c2_p1.font.size = Pt(14)
    c2_p1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    c2_p2 = c2_tf.add_paragraph()
    c2_p2.text = "$4.5 Million"
    c2_p2.font.size = Pt(26)
    c2_p2.font.bold = True
    c2_p2.font.color.rgb = RGBColor(0x0F, 0x20, 0x27)

    # Card 3: Headcount
    card3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.8), Inches(3.4), Inches(2.2))
    card3.fill.solid()
    card3.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF7)
    c3_tf = card3.text_frame
    c3_p1 = c3_tf.paragraphs[0]
    c3_p1.text = "Authorized FTE Target"
    c3_p1.font.size = Pt(14)
    c3_p1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    c3_p2 = c3_tf.add_paragraph()
    c3_p2.text = "85 Headcount FTEs"
    c3_p2.font.size = Pt(26)
    c3_p2.font.bold = True
    c3_p2.font.color.rgb = RGBColor(0x0F, 0x20, 0x27)

    # Slide 3: Table with different column order: Q1 Actual comes BEFORE Planned Spend!
    s3 = prs.slides.add_slide(blank_layout)
    t_box3 = s3.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11), Inches(0.8))
    p3 = t_box3.text_frame.paragraphs[0]
    p3.text = "Operational Units Performance"
    p3.font.name = "Georgia"
    p3.font.size = Pt(26)
    p3.font.bold = True

    table_shape = s3.shapes.add_table(5, 4, Inches(1), Inches(1.8), Inches(11), Inches(3.5))
    tbl = table_shape.table

    # Notice: Q1 Outlay is column 1, Full Year Cap is column 2!
    table_data = [
        ["Cost Center / Dept", "Q1 Outlay to Date", "Full Year Planned", "Status"],
        ["Engineering", "$510,000", "$2,100,000", "Active"],
        ["Marketing & Sales", "$290,000", "$1,200,000", "Active"],
        ["Operations & Support", "$180,000", "$800,000", "Active"],
        ["Human Resources", "$95,000", "$400,000", "Active"],
    ]
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            p_cell = cell.text_frame.paragraphs[0]
            p_cell.font.name = "Georgia"
            p_cell.font.size = Pt(13)
            if r_idx == 0:
                p_cell.font.bold = True
                p_cell.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x34, 0x49, 0x5E)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Saved variant deck: {output_path}")


def create_mapping_json(output_path: str):
    """Generates the budget mapping JSON configuration."""
    mapping = {
        "title": "Corporate Budget Plan Export",
        "sheets": [
            {
                "name": "Budget_Summary",
                "styles": {
                    "title": {
                        "font": "Segoe UI",
                        "size": 14,
                        "bold": True,
                        "font_color": "#1F4E78",
                    },
                    "label": {
                        "font": "Segoe UI",
                        "size": 11,
                        "bold": True,
                        "font_color": "#333333",
                        "align": "left",
                    },
                    "header": {
                        "font": "Segoe UI",
                        "size": 11,
                        "bold": True,
                        "bg_color": "#2F5597",
                        "font_color": "#FFFFFF",
                        "align": "center",
                        "border": "thin",
                    },
                    "currency": {
                        "font": "Segoe UI",
                        "size": 11,
                        "num_format": "$#,##0",
                        "align": "right",
                        "border": "thin",
                    },
                    "integer": {
                        "font": "Segoe UI",
                        "size": 11,
                        "num_format": "#,##0",
                        "align": "center",
                        "border": "thin",
                    },
                    "text": {
                        "font": "Segoe UI",
                        "size": 11,
                        "align": "left",
                        "border": "thin",
                    },
                },
                "cells": [
                    {
                        "cell": "B2",
                        "key": "fiscal_year",
                        "type": "string",
                        "description": "Fiscal year identifier, e.g. FY2025",
                        "style": "text",
                        "label": "Fiscal Year:",
                        "label_cell": "A2",
                        "label_style": "label",
                    },
                    {
                        "cell": "B3",
                        "key": "total_budget",
                        "type": "number",
                        "description": "Total approved expenditure or authorized budget across company",
                        "style": "currency",
                        "label": "Total Approved Budget:",
                        "label_cell": "A3",
                        "label_style": "label",
                    },
                    {
                        "cell": "B4",
                        "key": "headcount",
                        "type": "integer",
                        "description": "Total approved headcount or FTE count target",
                        "style": "integer",
                        "label": "Target Headcount (FTEs):",
                        "label_cell": "A4",
                        "label_style": "label",
                    },
                ],
                "tables": [
                    {
                        "key": "departments",
                        "start_cell": "A7",
                        "description": "Departmental breakdown of planned spend and Q1 actuals",
                        "columns": [
                            {
                                "header": "Department",
                                "key": "dept_name",
                                "type": "string",
                                "style": "text",
                                "header_style": "header",
                                "width": 25,
                            },
                            {
                                "header": "Planned Budget",
                                "key": "planned",
                                "type": "number",
                                "style": "currency",
                                "header_style": "header",
                                "width": 20,
                            },
                            {
                                "header": "Q1 Actual",
                                "key": "actual",
                                "type": "number",
                                "style": "currency",
                                "header_style": "header",
                                "width": 20,
                            },
                            {
                                "header": "Remaining Budget",
                                "key": "remaining",
                                "type": "formula",
                                "formula": "={planned}-{actual}",
                                "style": "currency",
                                "header_style": "header",
                                "width": 22,
                            },
                        ],
                    }
                ],
            }
        ],
    }

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(mapping, f, indent=2)
    print(f"Saved mapping JSON: {output_path}")


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(base_dir, "test_assets")

    create_sample_deck_standard(os.path.join(assets_dir, "sample_budget_deck.pptx"))
    create_sample_deck_variant(os.path.join(assets_dir, "sample_budget_variant_deck.pptx"))
    create_mapping_json(os.path.join(assets_dir, "budget_mapping.json"))
    print("All test assets created successfully.")
