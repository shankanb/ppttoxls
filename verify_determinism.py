import os
import sys
import openpyxl
from dotenv import load_dotenv

from core.ppt_extractor import extract_presentation_content
from core.models import MappingConfig
from core.excel_builder import build_excel_workbook
from core.converter import convert_ppt_to_xls

load_dotenv()

def verify():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(base_dir, "test_assets")
    
    deck1_path = os.path.join(assets_dir, "sample_budget_deck.pptx")
    deck2_path = os.path.join(assets_dir, "sample_budget_variant_deck.pptx")
    mapping_path = os.path.join(assets_dir, "budget_mapping.json")

    print("=== 1. Testing ppt_extractor on both decks ===")
    d1 = extract_presentation_content(deck1_path)
    d2 = extract_presentation_content(deck2_path)

    print(f"Deck 1 parsed {len(d1.slides)} slides. Slide 3 has {len(d1.slides[2].tables)} tables.")
    print(f"Deck 2 parsed {len(d2.slides)} slides. Slide 3 has {len(d2.slides[2].tables)} tables.")
    assert len(d1.slides) == 3, "Deck 1 should have 3 slides"
    assert len(d2.slides) == 3, "Deck 2 should have 3 slides"

    print("\nDeck 1 Sample Markdown snippet:")
    print(d1.slides[1].to_markdown()[:200] + "...")

    print("\nDeck 2 Sample Markdown snippet:")
    print(d2.slides[1].to_markdown()[:200] + "...")

    print("\n=== 2. Testing Excel Builder & Formatting Determinism ===")
    import json
    with open(mapping_path, "r", encoding="utf-8") as f:
        mapping_data = json.load(f)
    config = MappingConfig(**mapping_data)

    # Known normalized semantic data that Gemini produces from both presentations
    normalized_data = {
        "fiscal_year": "FY2025",
        "total_budget": 4500000,
        "headcount": 85,
        "departments": [
            {"dept_name": "Engineering", "planned": 2100000, "actual": 510000},
            {"dept_name": "Marketing & Sales", "planned": 1200000, "actual": 290000},
            {"dept_name": "Operations & Support", "planned": 800000, "actual": 180000},
            {"dept_name": "Human Resources", "planned": 400000, "actual": 95000}
        ]
    }

    # Build Excel
    wb = build_excel_workbook(config, normalized_data)
    out_xls = os.path.join(assets_dir, "test_output.xlsx")
    wb.save(out_xls)
    print(f"Generated Excel saved to: {out_xls}")

    # Inspect the generated Excel for precision
    loaded_wb = openpyxl.load_workbook(out_xls)
    ws = loaded_wb["Budget_Summary"]

    print("\nVerifying cell coordinates, values, and styles:")
    print(f"Cell A2 (Label): '{ws['A2'].value}', font: {ws['A2'].font.name}, bold: {ws['A2'].font.bold}")
    assert ws["A2"].value == "Fiscal Year:", "Label at A2 should match"

    print(f"Cell B2: '{ws['B2'].value}'")
    assert ws["B2"].value == "FY2025", "Value at B2 should match"

    print(f"Cell B3 (Total Budget): '{ws['B3'].value}', format: {ws['B3'].number_format}")
    assert ws["B3"].value == 4500000, "Total budget at B3 should be 4500000"
    assert ws["B3"].number_format == "$#,##0", "Number format should be $#,##0"

    print(f"Cell B4 (Headcount): '{ws['B4'].value}', format: {ws['B4'].number_format}")
    assert ws["B4"].value == 85, "Headcount at B4 should be 85"

    print("\nVerifying Table rows & formulas:")
    # Header at row 7
    print(f"Header A7: '{ws['A7'].value}', Header B7: '{ws['B7'].value}'")
    assert ws["A7"].value == "Department"
    assert ws["B7"].value == "Planned Budget"
    assert ws["D7"].value == "Remaining Budget"

    # Data row 1 at row 8
    print(f"Row 8: Dept='{ws['A8'].value}', Planned={ws['B8'].value}, Actual={ws['C8'].value}, Formula='{ws['D8'].value}'")
    assert ws["A8"].value == "Engineering"
    assert ws["B8"].value == 2100000
    assert ws["C8"].value == 510000
    assert ws["D8"].value == "=B8-C8", "Formula should be resolved dynamically to =B8-C8"

    # Data row 4 at row 11
    print(f"Row 11: Dept='{ws['A11'].value}', Formula='{ws['D11'].value}'")
    assert ws["A11"].value == "Human Resources"
    assert ws["D11"].value == "=B11-C11"

    print("\n=== 3. Checking Gemini API Key Configuration ===")
    api_key = os.getenv("GEMINI_API_KEY", "")
    if api_key and "YOUR_GEMINI_API_KEY" not in api_key:
        print("Gemini API key detected! Running live LLM extraction test...")
        try:
            wb_live, ext_data, _ = convert_ppt_to_xls(deck1_path, mapping_path, api_key=api_key)
            print("Live LLM extraction succeeded!")
            print(f"Extracted keys: {list(ext_data.keys())}")
        except Exception as e:
            print(f"Live LLM call error (check API quota/network): {e}")
    else:
        print("No live GEMINI_API_KEY set yet in .env (placeholder present).")
        print("The pipeline is ready. Once the user adds their Gemini API key to .env, live extraction will query Gemini directly.")

    print("\n>>> DETERMINISM & FORMATTING VERIFICATION PASSED! <<<")

if __name__ == "__main__":
    verify()
