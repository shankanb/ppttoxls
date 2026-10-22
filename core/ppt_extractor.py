import io
from typing import List, Union, BinaryIO
from pptx import Presentation


class SlideContent:
    def __init__(self, slide_number: int, title: str = ""):
        self.slide_number = slide_number
        self.title = title
        self.text_blocks: List[str] = []
        self.tables: List[List[List[str]]] = []

    def to_markdown(self) -> str:
        lines = [f"=== SLIDE {self.slide_number}: {self.title or 'Untitled'} ==="]
        if self.text_blocks:
            lines.append("Content:")
            for block in self.text_blocks:
                lines.append(f"- {block}")
        if self.tables:
            lines.append("Tables:")
            for t_idx, table in enumerate(self.tables, 1):
                lines.append(f"Table {t_idx}:")
                if not table:
                    continue
                header = table[0]
                lines.append("| " + " | ".join(header) + " |")
                lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                for row in table[1:]:
                    lines.append("| " + " | ".join(row) + " |")
        return "\n".join(lines)


class DeckContent:
    def __init__(self):
        self.slides: List[SlideContent] = []

    def add_slide(self, slide: SlideContent):
        self.slides.append(slide)

    def to_markdown(self) -> str:
        return "\n\n".join(s.to_markdown() for s in self.slides)


def extract_presentation_content(file_source: Union[str, BinaryIO, bytes]) -> DeckContent:
    """
    Extracts structured text, titles, and tables from a PowerPoint presentation (.pptx).
    Accepts a filepath, file-like binary stream (e.g. Django uploaded file), or raw bytes.
    """
    if isinstance(file_source, bytes):
        file_source = io.BytesIO(file_source)

    prs = Presentation(file_source)
    deck = DeckContent()

    for idx, slide in enumerate(prs.slides, start=1):
        # Extract title
        title = ""
        try:
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass

        slide_data = SlideContent(slide_number=idx, title=title)

        for shape in slide.shapes:
            # Skip title shape if already captured as title
            if title and shape == getattr(slide.shapes, "title", None):
                continue

            # 1. Check for Tables
            if shape.has_table:
                table_matrix: List[List[str]] = []
                for row in shape.table.rows:
                    row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_matrix.append(row_data)
                if table_matrix:
                    slide_data.tables.append(table_matrix)

            # 2. Check for Text Frames
            elif shape.has_text_frame:
                text_frame_text = shape.text_frame.text.strip()
                if text_frame_text:
                    # Clean up multiple newlines
                    cleaned = "\n".join(
                        p.strip() for p in text_frame_text.splitlines() if p.strip()
                    )
                    if cleaned:
                        slide_data.text_blocks.append(cleaned)

        deck.add_slide(slide_data)

    return deck
